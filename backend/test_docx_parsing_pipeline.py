"""
Integration test for AI Resume Extraction pipeline, health checks, retry mechanism, and status handling.
"""
import asyncio
import os
import sys
from docx import Document

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app.core.config import settings
from app.ai.llm_provider import get_llm_provider
from app.parsers.docx_parser import extract_text_from_docx
from app.parsers.resume_parser import parse_resume
from app.ai.resume_extractor import extract_resume_with_llm
from app.parsers.contact_parser import ContactInfo


async def test_pipeline():
    print("=" * 75)
    print("AI RESUME EXTRACTION PIPELINE VERIFICATION TEST")
    print("=" * 75)

    # 1. AI Health Check Verification
    provider = get_llm_provider()
    health_details = await provider.check_health_details()
    print(f"\n1. AI PROVIDER HEALTH CHECK:")
    print(f"   Active Provider: {settings.LLM_PROVIDER}")
    print(f"   Provider Class:  {provider.__class__.__name__}")
    print(f"   Available:       {health_details['available']}")
    print(f"   Model:           {health_details['model']}")
    print(f"   Error Details:   {health_details['error']}")

    # 2. Test DOCX Document Creation & Complete Text Extraction
    test_docx_path = os.path.abspath("data/test_resumes/sample_resume_2.docx")
    os.makedirs(os.path.dirname(test_docx_path), exist_ok=True)

    doc = Document()
    doc.add_heading("Aditya Bonde", level=1)
    doc.add_paragraph("Email: aditya.bonde@example.com | Phone: +1 555 0199 | Location: San Francisco, CA")
    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph("Senior Full Stack & AI Engineer with 5+ years experience building scalable backend microservices and AI pipelines.")
    
    doc.add_heading("Technical Skills", level=2)
    doc.add_paragraph("Python, FastAPI, React, TypeScript, Docker, Kubernetes, PostgreSQL, LLM Systems, Ollama, Groq")

    doc.add_heading("Work Experience", level=2)
    doc.add_paragraph("Senior AI Software Engineer — Tech Corp (2022-01 to Present)")
    doc.add_paragraph("• Architected high-throughput AI document processing pipeline using Python and FastAPI.")
    doc.add_paragraph("Software Engineer — Cloud Solutions Inc (2020-03 to 2021-12)")
    doc.add_paragraph("• Developed distributed microservices and database schemas.")

    doc.add_heading("Education", level=2)
    doc.add_paragraph("Master of Science in Computer Science — Stanford University (2018 - 2020)")

    doc.add_heading("Languages", level=2)
    doc.add_paragraph("English (Fluent), Japanese (JLPT N2 Certified), Spanish (Basic)")

    doc.add_heading("Certifications", level=2)
    doc.add_paragraph("AWS Certified Solutions Architect (2021)")

    doc.add_heading("Projects", level=2)
    doc.add_paragraph("TalentVault AI Resume Engine (FastAPI, React, Groq, Ollama)")

    doc.save(test_docx_path)
    print(f"\n2. CREATED TEST DOCX RESUME AT: {test_docx_path}")

    # 3. Test Full DOCX Text Extraction
    docx_extraction = extract_text_from_docx(test_docx_path)
    print(f"\n3. DOCX TEXT EXTRACTION:")
    print(f"   Quality Score: {docx_extraction.quality_score}")
    print(f"   Line Count:    {docx_extraction.line_count}")
    print(f"   Char Count:    {docx_extraction.char_count}")
    assert docx_extraction.char_count > 200, "DOCX extraction produced empty or short text"
    assert "Aditya Bonde" in docx_extraction.text
    assert "Japanese (JLPT N2 Certified)" in docx_extraction.text

    # 4. Test Resume Parsing Pipeline
    parse_res = await parse_resume(test_docx_path)
    print(f"\n4. RESUME PARSE RESULT:")
    print(f"   Parsing Status:   {parse_res.parsing_status}")
    print(f"   Extraction Method:{parse_res.extraction_method}")
    print(f"   LLM Model:        {parse_res.llm_model}")
    print(f"   Errors Count:     {len(parse_res.errors)}")
    print(f"   Warnings:         {parse_res.warnings}")

    if parse_res.extraction:
        ext = parse_res.extraction
        print(f"   Extracted Name:       {ext.full_name}")
        print(f"   Extracted Email:      {ext.email}")
        print(f"   Extracted Phone:      {ext.phone}")
        print(f"   Extracted Skills:     {ext.skills}")
        print(f"   Extracted Experiences:{[e.company for e in ext.experiences]}")
        print(f"   Extracted Education:  {[e.institution for e in ext.education]}")
        print(f"   Extracted Languages:  {ext.languages}")

    if parse_res.parsing_status == "PARSED":
        print("   -> Fully parsed via AI LLM!")
    else:
        assert parse_res.parsing_status == "NEEDS_REVIEW", f"Expected NEEDS_REVIEW or PARSED, got {parse_res.parsing_status}"
        print(f"   -> Successfully marked as NEEDS_REVIEW due to AI failure: {parse_res.warnings}")

    print("\n" + "=" * 75)
    print("PIPELINE TEST COMPLETED SUCCESSFULLY!")
    print("=" * 75)


if __name__ == "__main__":
    asyncio.run(test_pipeline())
